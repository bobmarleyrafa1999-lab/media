#!/usr/bin/env python3
"""Editable Word version of a Moodius journal (opens in Word, Pages and Google Docs).

    python3 make_docx.py [content/the-unexamined-self.yaml] [out/The-Unexamined-Self-editable.docx]

The YAML stays the source of truth; this is for editing or proofreading in a word processor.
"""
import os, sys, yaml
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SLATE, MUTED, EMBER = RGBColor(0x1B, 0x1F, 0x24), RGBColor(0x59, 0x60, 0x6A), RGBColor(0x98, 0x52, 0x1B)
ROMAN = ["", "I", "II", "III", "IV", "V", "VI"]


def main(src, out):
    d = yaml.safe_load(open(src, encoding="utf-8"))
    doc = Document()
    st = doc.styles
    st["Normal"].font.name = "Inter"; st["Normal"].font.size = Pt(10.5); st["Normal"].font.color.rgb = SLATE
    for name, size in [("Title", 30), ("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 12)]:
        st[name].font.name = "Playfair Display"; st[name].font.size = Pt(size); st[name].font.color.rgb = SLATE

    def kicker(text):
        r = doc.add_paragraph().add_run(text.upper()); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = EMBER

    def quote(q):
        p = doc.add_paragraph(); r = p.add_run(q["text"]); r.italic = True; r.font.size = Pt(12)
        c = doc.add_paragraph().add_run(f"{q['cite']} · {'verbatim' if q['kind'] == 'verbatim' else 'paraphrase'}")
        c.font.size = Pt(8); c.font.color.rgb = MUTED

    def lines(n=6):
        for _ in range(n):
            doc.add_paragraph("_" * 70).runs[0].font.color.rgb = RGBColor(0xC8, 0xC2, 0xB6)

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    m = d["meta"]
    doc.add_paragraph(m["title"], style="Title")
    doc.add_paragraph(m["subtitle"])
    if d.get("epigraph"):
        quote(d["epigraph"])
    page_break()
    for f in d.get("front", []):
        doc.add_heading(f["heading"], 1)
        for para in f["body"].strip().split("\n\n"):
            doc.add_paragraph(" ".join(para.split()))
        page_break()
    for part in d["parts"]:
        kicker(f"Part {ROMAN[part['n']]}")
        doc.add_heading(part["title"], 1)
        doc.add_paragraph(part["intro"])
        page_break()
        for day in [x for x in d["days"] if x["part"] == part["n"]]:
            kicker(f"Day {day['n']} · Part {ROMAN[part['n']]} · {part['title']}")
            doc.add_heading(day["title"], 2)
            quote(day["quote"])
            for para in day["reading"].strip().split("\n\n"):
                doc.add_paragraph(" ".join(para.split()))
            ev = day["evidence"]
            p = doc.add_paragraph(); r = p.add_run(f"Evidence · {ev['tier']}: "); r.bold = True; r.font.color.rgb = EMBER
            p.add_run(ev["note"]).font.size = Pt(9)
            kicker("Questions")
            for i, q in enumerate(day["prompts"], 1):
                doc.add_paragraph(f"{i}. {q}"); lines(4)
            kicker("Exercise")
            doc.add_heading(day["exercise"]["name"], 3)
            doc.add_paragraph(" ".join(day["exercise"]["steps"].split()))
            lines(8)
            kicker("Tonight")
            doc.add_paragraph(day["tonight"]); lines(2)
            page_break()
        if d.get("review"):
            kicker(f"{d['review']['heading']} · Part {ROMAN[part['n']]}")
            for q in d["review"]["prompts"] + [d["review"]["scale"]]:
                doc.add_paragraph(q); lines(3)
            page_break()
    if d.get("vocabulary"):
        v = d["vocabulary"]; doc.add_heading(v["heading"], 1); doc.add_paragraph(v["intro"])
        for g in v["groups"]:
            p = doc.add_paragraph(); p.add_run(g["name"] + ": ").bold = True; p.add_run(", ".join(g["words"]))
    if d.get("values"):
        v = d["values"]; doc.add_heading(v["heading"], 1); doc.add_paragraph(v["intro"]); doc.add_paragraph(", ".join(v["words"]))
    if d.get("references"):
        doc.add_heading("References", 1)
        for k in sorted(d["references"], key=lambda k: d["references"][k].lower()):
            doc.add_paragraph(d["references"][k]).runs[0].font.size = Pt(9)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "content", "the-unexamined-self.yaml")
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "out", "The-Unexamined-Self-editable.docx")
    main(src, out)
