#!/usr/bin/env python3
"""Build the written deliverables from concepts.py, sequences.py, questions.py and the calendar below.

Writes (into the parent folder): 02-concepts.md/.csv, 03-sequences.md, 05-questions.md/.csv,
07-calendar.md/.csv, 09-tracker.xlsx, 09-tracker.csv, data.json (for the overview page).
"""
import csv, json, os, sys, datetime as dt
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from concepts import as_dicts, PILLARS, FORMATS
from sequences import S
from questions import rows as qrows

CON = {c["id"]: c for c in as_dicts()}
SEQ = {s["id"]: s for s in S}

# 30 days, 8 Nov - 7 Dec 2026. Each day: 21:00 evening set; its payoff runs 09:00 the next morning.
# Mon = about-you question (rule 6); Fri = you pick (rule 7). hand = suggested night to hand-post the real sticker (test X11).
CAL = [
 ("2026-11-08", "S01", False), ("2026-11-09", "C28", False), ("2026-11-10", "S03", True),
 ("2026-11-11", "S07", False), ("2026-11-12", "S02", False), ("2026-11-13", "X03", True),
 ("2026-11-14", "C40", True), ("2026-11-15", "S08", False), ("2026-11-16", "S12", False),
 ("2026-11-17", "S04", True), ("2026-11-18", "C45", False), ("2026-11-19", "S10", False),
 ("2026-11-20", "X03", True), ("2026-11-21", "S06", True), ("2026-11-22", "S15", False),
 ("2026-11-23", "C35", False), ("2026-11-24", "S09", True), ("2026-11-25", "S14", False),
 ("2026-11-26", "C41", False), ("2026-11-27", "X03", True), ("2026-11-28", "S11", True),
 ("2026-11-29", "C20", False), ("2026-11-30", "S05", False), ("2026-12-01", "S13", True),
 ("2026-12-02", "C42", False), ("2026-12-03", "C19", False), ("2026-12-04", "X03", True),
 ("2026-12-05", "X02", True), ("2026-12-06", "C23", False), ("2026-12-07", "C37", False),
]
NOTES = {"2026-11-12": "Epictetus arc week (12-15 Nov)", "2026-11-14": "Epictetus arc week",
         "2026-11-19": "19 Nov review: read tracker first", "2026-11-28": "Answers feed next week's Reading Room carousel (with consent)"}


def cal_rows():
    out = []
    for date, ref, hand in CAL:
        d = dt.date.fromisoformat(date)
        nxt = (d + dt.timedelta(days=1)).isoformat()
        if ref in SEQ:
            s = SEQ[ref]; c = CON[s["concept"]]
            title, pillar, fmt = s["name"], s["pillar"], s["format"]
            ev = len(s["evening"]); mo = len(s["morning"])
            hook = s["evening"][0]["title"]; kind = "sequence (rendered)"
        else:
            c = CON[ref]; title, pillar, fmt = c["hook"][:60], c["pillar"], c["format"]
            ev, mo, hook, kind = 2, 2, c["hook"], "concept (routine renders)"
        out.append(dict(date=date, day=d.strftime("%a"), evening_2100=f"{ref}: {title}", frames_evening=ev,
                        morning_0900=nxt, frames_morning=mo, ref=ref, kind=kind, pillar=PILLARS[pillar],
                        format=FORMATS[fmt], hook=hook, hand_sticker_night="yes" if hand else "",
                        note=NOTES.get(date, "")))
    return out


def md_escape(s):
    return (s or "").replace("|", "/").replace("\n", " ")


def write_concepts():
    cs = [c for c in as_dicts()]
    with open(os.path.join(ROOT, "02-concepts.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(cs[0].keys())); w.writeheader(); w.writerows(cs)
    L = ["# 50 Story concepts (+3 recurring formats)", "",
         "Each concept: hook, layout (template), the sticker to add by hand, the reply prompt for the scheduled version, a discussion question, the follow-up story, an optional CTA and the source. The 15 strongest are written out frame by frame in 03-sequences.md.", "",
         "Moodius takes are drafts in the page's voice: change anything that isn't true for you.", ""]
    for fmt, name in FORMATS.items():
        group = [c for c in cs if c["format"] == fmt]
        if not group: continue
        L += [f"## {name}", ""]
        for c in group:
            L += [f"**{c['id']} · {c['pillar_name']}**: {c['hook']}", "",
                  f"- Layout: {c['layout']}", f"- Sticker (by hand): {c['sticker']}", f"- Scheduled version: {c['reply_fallback']}",
                  f"- Discussion question: {c['question']}", f"- Follow-up story: {c['follow_up']}"]
            if c["cta"]: L.append(f"- CTA: {c['cta']}")
            L += [f"- Source: {c['source']}", ""]
    open(os.path.join(ROOT, "02-concepts.md"), "w").write("\n".join(L))


def write_sequences(index):
    L = ["# The 15 strongest Story sequences, ready to publish", "",
         "Evening set at 21:00 (hook + sticker), morning set at 09:00 the next day (payoff + Moodius take).",
         "Images: renders/<id>/. `_auto` = reply prompt printed, safe to schedule in Metricool. `_hand` = clear space for the real sticker; post it from the Instagram app and add the sticker listed. `_all` = no sticker on that frame.", "",
         "Frames in [square brackets] are filled by hand (real answers or your own pick); never invent them.", ""]
    for s in S:
        c = CON[s["concept"]]
        L += [f"## {s['id']} · {s['name']} ({PILLARS[s['pillar']]}, {FORMATS[s['format']]})", "", f"*Why it's strong:* {s['why']}", "", f"Source: {c['source']}", ""]
        n = 0
        for part, frames in (("Evening, 21:00", s["evening"]), ("Next morning, 09:00", s["morning"])):
            L += [f"**{part}**", ""]
            for f in frames:
                n += 1
                bits = [x for x in [f.get("kicker") and f"*{f['kicker']}*", f.get("title") and f"**{f['title']}**", f.get("body"),
                                    f.get("left") and f"A: {f['left']} / B: {f['right']}" if f.get("left") else None,
                                    f.get("options") and " / ".join(f"{'ABCD'[i]} {o}" for i, o in enumerate(f["options"])),
                                    f.get("note") and f"_{f['note']}_", f.get("footer") and f"CTA: {f['footer']}"] if x]
                L.append(f"{n}. [{f['t']}] " + " · ".join(bits))
                if f.get("sticker"): L.append(f"   - By hand: {f['sticker']}")
                if f.get("reply"): L.append(f"   - Scheduled version says: “{f['reply']}”")
            L.append("")
        files = [i["file"] for i in index if i["seq"] == s["id"]]
        L += ["Files: " + ", ".join(f"`{x}`" for x in files), ""]
    open(os.path.join(ROOT, "03-sequences.md"), "w").write("\n".join(L))


def write_questions():
    qs = list(qrows())
    with open(os.path.join(ROOT, "05-questions.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(qs[0].keys())); w.writeheader(); w.writerows(qs)
    L = ["# 100 questions that start conversations", "",
         "20 per pillar. Sticker = the one to add by hand; the scheduled version uses the reply line. Depth: light (one tap), medium (one line), deep (a real answer: max 2 a week, and Moodius answers first).", ""]
    for p, name in PILLARS.items():
        L += [f"## {name}", "", "| # | Question | Sticker | Depth |", "|---|---|---|---|"]
        L += [f"| {q['id']} | {md_escape(q['question'])} | {q['sticker']} | {q['depth']} |" for q in qs if q["pillar"] == p]
        L.append("")
    open(os.path.join(ROOT, "05-questions.md"), "w").write("\n".join(L))
    return qs


def write_calendar():
    rs = cal_rows()
    with open(os.path.join(ROOT, "07-calendar.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rs[0].keys())); w.writeheader(); w.writerows(rs)
    L = ["# 30-day Stories calendar: Sun 8 Nov to Mon 7 Dec 2026", "",
         "26 Oct to 7 Nov is already scheduled (22 quiz stories from the engagement thread). This picks up the next day.",
         "Two story slots a day as now: 21:00 evening set, 09:00 payoff next morning. Monday nights ask about you; Friday nights are 'you pick'.",
         "'Hand sticker night' = the nights to post the real sticker version from your phone (test X11). Skipping it breaks nothing: the scheduled reply version runs anyway.", "",
         "| Date | Day | 21:00 evening | Pillar | Format | Hand sticker night | Note |", "|---|---|---|---|---|---|---|"]
    for r in rs:
        L.append(f"| {r['date']} | {r['day']} | {md_escape(r['evening_2100'])} | {r['pillar']} | {r['format']} | {r['hand_sticker_night']} | {r['note']} |")
    from collections import Counter
    cnt = Counter(r["pillar"] for r in rs)
    L += ["", "Pillar mix over 30 nights: " + ", ".join(f"{k} {v}" for k, v in cnt.most_common()) + "."]
    open(os.path.join(ROOT, "07-calendar.md"), "w").write("\n".join(L))
    return rs


def write_tracker(cal):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    head = Font(bold=True, color="F3EFE7"); fill = PatternFill("solid", fgColor="1B1F24")
    # Sheet 1: per story frame
    ws = wb.active; ws.title = "Story log"
    cols = ["date", "slot", "ref", "frame", "version (auto/hand)", "sticker", "reach", "views", "taps forward", "taps back", "exits",
            "next story", "replies (count in DMs)", "sticker taps (votes/answers)", "shares", "profile visits", "follows",
            "exit rate", "replies per 100 reached", "votes per 100 reached", "notes"]
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        ws.cell(1, i).font = head; ws.cell(1, i).fill = fill; ws.cell(1, i).alignment = Alignment(wrap_text=True)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 14
    r = 2
    for row in cal:
        ref = row["ref"]
        if ref in SEQ:
            s = SEQ[ref]; frames = [("21:00", f) for f in s["evening"]] + [("09:00", f) for f in s["morning"]]
        else:
            frames = [("21:00", {"sticker": CON[ref]["sticker"]}), ("21:00", {}), ("09:00", {}), ("09:00", {})]
        for k, (slot, f) in enumerate(frames, 1):
            day = row["date"] if slot == "21:00" else row["morning_0900"]
            ws.append([day, slot, ref, k, "", (f.get("sticker") or "")[:40]])
            ws[f"R{r}"] = f'=IF(G{r}>0,K{r}/G{r},"")'; ws[f"S{r}"] = f'=IF(G{r}>0,M{r}/G{r}*100,"")'; ws[f"T{r}"] = f'=IF(G{r}>0,N{r}/G{r}*100,"")'
            for col in "RS T".replace(" ", ""):
                ws[f"{col}{r}"].number_format = "0.0" if col != "R" else "0%"
            r += 1
    dv = DataValidation(type="list", formula1='"auto,hand"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"E2:E{r}")
    ws.freeze_panes = "A2"
    # Sheet 2: per sequence summary (completion = last frame reach / first frame reach)
    s2 = wb.create_sheet("Sequence summary")
    hdr = ["ref", "first-frame reach", "last-frame reach", "completion", "total replies", "replies per 100 reached (first frame)",
           "total sticker taps", "shares", "follows", "follows per 1,000 reached", "verdict (keep / tweak / drop)"]
    s2.append(hdr)
    for i, c in enumerate(hdr, 1):
        s2.cell(1, i).font = head; s2.cell(1, i).fill = fill; s2.cell(1, i).alignment = Alignment(wrap_text=True)
        s2.column_dimensions[s2.cell(1, i).column_letter].width = 16
    for j, row in enumerate(cal, 2):
        ref = row["ref"]; L = "'Story log'!"
        s2.append([ref])
        s2[f"B{j}"] = f'=IFERROR(INDEX({L}G:G,MATCH("{ref}",{L}C:C,0)),"")'
        s2[f"C{j}"] = f'=IFERROR(LOOKUP(2,1/({L}C2:C500="{ref}"),{L}G2:G500),"")'
        s2[f"D{j}"] = f'=IF(AND(ISNUMBER(B{j}),B{j}>0,ISNUMBER(C{j})),C{j}/B{j},"")'; s2[f"D{j}"].number_format = "0%"
        s2[f"E{j}"] = f'=SUMIF({L}C:C,"{ref}",{L}M:M)'
        s2[f"F{j}"] = f'=IF(AND(ISNUMBER(B{j}),B{j}>0),E{j}/B{j}*100,"")'; s2[f"F{j}"].number_format = "0.0"
        s2[f"G{j}"] = f'=SUMIF({L}C:C,"{ref}",{L}N:N)'
        s2[f"H{j}"] = f'=SUMIF({L}C:C,"{ref}",{L}O:O)'
        s2[f"I{j}"] = f'=SUMIF({L}C:C,"{ref}",{L}Q:Q)'
        s2[f"J{j}"] = f'=IF(AND(ISNUMBER(B{j}),B{j}>0),I{j}/B{j}*1000,"")'; s2[f"J{j}"].number_format = "0.0"
    # Sheet 3: weekly account numbers
    s3 = wb.create_sheet("Weekly")
    s3.append(["week starting", "followers (Sunday)", "new followers", "accounts reached (7 days)", "story reach, average per story",
               "replies (all stories)", "DM conversations started", "profile visits", "what worked", "what to change"])
    for i in range(1, 11):
        s3.cell(1, i).font = head; s3.cell(1, i).fill = fill; s3.column_dimensions[s3.cell(1, i).column_letter].width = 18
    for k in range(5):
        s3.append([(dt.date(2026, 11, 8) + dt.timedelta(days=7 * k)).isoformat()])
    # Sheet 4: definitions
    s4 = wb.create_sheet("How to fill")
    for line in [
        "Where the numbers come from: Instagram app > your story (or Archive > Stories) > swipe up > Insights. Fill within 24 hours, or from the archive later.",
        "Replies: Instagram's API returns 0 for story replies from accounts in Europe, so count them in your DMs (or the 'Replies' line in the app).",
        "Sticker taps: the poll / quiz / slider / question results on the hand-posted version (swipe up on the story).",
        "Completion = reach of the last frame / reach of the first frame of the same sequence.",
        "Exit rate = exits / reach for that frame. A frame with a much higher exit rate than its neighbours is the weak one.",
        "Replies per 100 reached is the main score for stories (Instagram ranks stories partly on replies). Follows per 1,000 reached is the growth score.",
        "Small numbers: with under 100 viewers a story, one reply moves the rate a lot. Judge on 5+ stories per arm, never on one night.",
        "Instagram does not report story metrics below 5 for some fields; leave the cell empty, don't write 0.",
    ]:
        s4.append([line])
    s4.column_dimensions["A"].width = 140
    wb.save(os.path.join(ROOT, "09-tracker.xlsx"))
    with open(os.path.join(ROOT, "09-tracker.csv"), "w", newline="") as f:
        csv.writer(f).writerow(cols[:17] + ["notes"])


def main():
    idx = json.load(open(os.path.join(ROOT, "renders", "index.json")))
    write_concepts(); write_sequences(idx); qs = write_questions(); cal = write_calendar(); write_tracker(cal)
    data = dict(concepts=as_dicts(), sequences=S, questions=qs, calendar=cal, renders=idx, pillars=PILLARS, formats=FORMATS)
    json.dump(data, open(os.path.join(ROOT, "data.json"), "w"), ensure_ascii=False)
    print("ok", len(as_dicts()), len(S), len(qs), len(cal))


if __name__ == "__main__":
    main()
